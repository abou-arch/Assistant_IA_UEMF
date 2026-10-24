# Recrée base_QR_EIDIA.jsonl à partir du classeur Excel, après relecture ou ajout de questions.
# Utilisation : python export_jsonl.py base_QR_EIDIA.xlsx
import json
import sys

from openpyxl import load_workbook

LIGNE_ENTETE = 4

nom_classeur = sys.argv[1] if len(sys.argv) > 1 else "base_QR_EIDIA.xlsx"
nom_sortie = nom_classeur.replace(".xlsx", ".jsonl")

classeur = load_workbook(nom_classeur, data_only=True)


def lire_feuille(nom_feuille):
    """Renvoie une liste de dictionnaires {titre de colonne: valeur}, une entrée par ligne remplie."""
    feuille = classeur[nom_feuille]
    titres = [cellule.value for cellule in feuille[LIGNE_ENTETE]]
    lignes = []
    for valeurs in feuille.iter_rows(min_row=LIGNE_ENTETE + 1, values_only=True):
        if valeurs[0] is None:
            continue
        lignes.append(dict(zip(titres, valeurs)))
    return lignes


source_par_id = {}
for source in lire_feuille("Sources"):
    source_par_id[source["ID"]] = {"id": source["ID"], "titre": source["Titre"], "url": source["URL"] or ""}

nombre_ecrites = 0
with open(nom_sortie, "w", encoding="utf-8") as fichier:
    for ligne in lire_feuille("Questions-Réponses"):
        identifiants_sources = [texte.strip() for texte in str(ligne["Sources"] or "").split(",") if texte.strip()]
        enregistrement = {
            "id": ligne["ID"],
            "categorie": ligne["Catégorie"],
            "question": ligne["Question (FR)"],
            "variante": ligne["Variante (FR)"],
            "question_en": ligne["Question (EN)"],
            "question_ar": ligne["Question (AR)"],
            "reponse": ligne["Réponse attendue"],
            "comportement_attendu": ligne["Comportement attendu"],
            "orientation": ligne["Orienter vers"] or "",
            "sources": [source_par_id[identifiant] for identifiant in identifiants_sources if identifiant in source_par_id],
            "validite": ligne["Validité"],
            "origine": ligne["Origine"],
            "type_question": ligne["Type de question"],
            "statut": ligne["Statut"],
            "usage": ligne["Usage"],
            "note_verification": ligne["Note de vérification"] or "",
            "relu_par": ligne["Relu par"] or "",
            "a_indexer": ligne["Usage"] == "Base + évaluation",
            "texte_a_indexer": "Question : " + ligne["Question (FR)"] + "\nRéponse : " + ligne["Réponse attendue"],
        }
        fichier.write(json.dumps(enregistrement, ensure_ascii=False) + "\n")
        nombre_ecrites += 1

print(nombre_ecrites, "questions exportées dans", nom_sortie)
